"""Build the patent-coverage snapshot for the LAPaL prototype.

Usage (from the project folder):
    python3 tools/build_patents.py            # use cached downloads where present
    python3 tools/build_patents.py --refresh  # re-download LAPaL pages and MedsPaL JSONs

Inputs
    URL_review_to_confirm.xlsx          Lobna's confirmed row -> lapal.ch URL list
    LAPAL_Patent info_MASTER-copy.xlsx  Follow-up tab: TA, category, MedsPaL uuid
    lapal.ch entry pages                patent cards, and the embedded MedsPaL uuid
    assets.medspal.org product + patent JSONs
    World Bank country API              names, ISO codes, income level

Outputs (data/)
    patents_data.json        the snapshot the prototype reads
    build_report.md          what was included, excluded, estimated or not understood
    card_status_review.xlsx  every LAPaL card status line with how it was read, for checking

Data assumptions, all reported rather than resolved silently:
- Rows cleared in the review sheet (no product name left) are excluded; a blank CONFIRM is
  treated as accepted. Lobna left CONFIRM blank and cleared the rows she rejected.
- Rows pointing to the same page become one entry.
- An entry uses MedsPaL data when its LAPaL page embeds a MedsPaL uuid. If the page has none but
  the workbook row has one, that uuid is used and flagged. Otherwise the LAPaL cards are used and
  the entry is marked preliminary.
- Three map classes: granted / filed / not_in_force. Rejected counts as not in force (Lobna's call).
  Statuses that cannot be placed (e.g. "To be confirmed") are kept but not mapped.
- A regional office named without its member list is expanded to every member (Lobna's call).
- Card families have no filing date on the card: it is looked up on Google Patents from the
  representative patent (cached in cache/google_patents_filing.json); if that fails we estimate
  expiry minus 20 years and flag it.
"""
import argparse, datetime as dt, html, json, os, re, sys, time, urllib.request
from collections import Counter, defaultdict
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CACHE = os.path.join(ROOT, 'cache')
OUT = os.path.join(ROOT, 'data')
# Licence territories, read-only, from the licence-territory dashboard's workbook (same sheets the
# dashboard reads: CAB from 'CAB-LA master', LEN and alimatravir from 'Master comparison').
LICENCE_XLSX = os.path.join(os.path.dirname(ROOT), 'licence-territory-dashboard', 'data', 'HIV_PrEP_licence_territories_v2_3.xlsx')
# LAPaL entries a voluntary licence applies to (Lobna, 27 Oct). Combinations and other formulations
# (Q4M, pacfosacil, CAB + LEN, ISL + LEN) are left out until confirmed.
LICENCES = [
    {'code': 'CAB', 'label': 'Cabotegravir', 'licensor': 'ViiV Healthcare', 'type': 'MPP licence',
     'sheet': 'CAB-LA master', 'iso3_col': 1, 'status_col': 4, 'basis_col': 5,
     'entries': ['compounds/cabotegravir', 'formulations/cab-la-cabotegravir-2-monthly']},
    {'code': 'LEN', 'label': 'Lenacapavir', 'licensor': 'Gilead Sciences', 'type': 'bilateral licence',
     'sheet': 'Master comparison', 'iso3_col': 1, 'status_col': 5, 'basis_col': None,
     'entries': ['compounds/lenacapavir-len', 'formulations/lenacapavir-6-monthly', 'formulations/lenacapavir-once-yearly']},
    {'code': 'AMI', 'label': 'Alimatravir', 'licensor': 'MSD', 'type': 'bilateral licence',
     'sheet': 'Master comparison', 'iso3_col': 1, 'status_col': 6, 'basis_col': None,
     'entries': ['compounds/alimatravir-mk-8527']},
]
REVIEW = os.path.join(ROOT, 'URL_review_to_confirm.xlsx')
MASTER = os.path.join(ROOT, 'LAPAL_Patent info_MASTER-copy.xlsx')

# --- vocabularies ------------------------------------------------------------------------------

STATUS_CLASS = {
    # Granted: in force, including while opposed, after amendment, or after restoration of a
    # lapsed right; an extension request does not change that the patent is granted.
    'granted': 'granted', 'term extended': 'granted', 'granted (opposed)': 'granted',
    'granted (claims amended)': 'granted', 'restored': 'granted', 'term extensions filed': 'granted',
    'filed': 'filed', 'filed (opposed)': 'filed', 'pending': 'filed',
    'not in force': 'not_in_force', 'withdrawn': 'not_in_force', 'rejected': 'not_in_force',
    'rejected - under appeal': 'not_in_force', 'expired': 'not_in_force', 'lapsed': 'not_in_force',
    'ceased': 'not_in_force', 'abandoned': 'not_in_force', 'revoked': 'not_in_force',
}
RANK = {'granted': 3, 'filed': 2, 'not_in_force': 1}

# Regional patent offices. When a card names an office with its own member list in brackets we
# use that list; when it names the office alone we expand to all members below.
REGIONS = {
    'EP': {'label': 'European Patent Office (EPO), 39 member states',
           'members': 'AL AT BE BG CH CY CZ DE DK EE ES FI FR GB GR HR HU IE IS IT LI LT LU LV MC ME MK MT NL NO PL PT RO RS SE SI SK SM TR'.split()},
    'EA': {'label': 'Eurasian Patent Organization (EAPO)',
           'members': 'AM AZ BY KG KZ RU TJ TM'.split()},
    # The 18 Harare Protocol states as LAPaL curators list them in "AP (...)" (Cabo Verde, which
    # joined in 2022, does not appear in any card, so it is left out until confirmed).
    'AP': {'label': 'African Regional Intellectual Property Organization (ARIPO), Harare Protocol states',
           'members': 'BW GH GM KE LR LS MW MZ NA RW SD SL ST SZ TZ UG ZM ZW'.split()},
    'OA': {'label': 'African Intellectual Property Organization (OAPI)',
           'members': 'BF BJ CF CG CI CM GA GN GQ GW KM ML MR NE SN TD TG'.split()},
    'GC': {'label': 'Gulf Cooperation Council Patent Office (GCCPO)',
           'members': 'AE BH KW OM QA SA'.split()},
}
REGION_ALIAS = {'EP': 'EP', 'EPO': 'EP', 'EPC': 'EP', 'EUROPE': 'EP', 'EUROPEAN PATENT OFFICE': 'EP',
                'EPO MEMBER STATES': 'EP', 'EA': 'EA', 'EAPO': 'EA', 'AP': 'AP', 'ARIPO': 'AP',
                'OA': 'OA', 'OAPI': 'OA', 'GC': 'GC', 'GCC': 'GC'}
CODE_ALIAS = {'UK': 'GB', 'US': 'US', 'USA': 'US'}
MEDSPAL_CODE_ALIAS = {'KV': 'XK'}  # MedsPaL writes Kosovo as KV; the World Bank uses XK
# Former states on old records (Yugoslavia, Serbia and Montenegro, Soviet Union): not mappable today.
HISTORICAL_CODES = {'YU', 'CS', 'SU'}
NAME_ALIAS = {'UNITED STATES': 'US', 'UNITED STATES OF AMERICA': 'US', 'SOUTH KOREA': 'KR',
              'REPUBLIC OF KOREA': 'KR', 'KOREA': 'KR', 'RUSSIA': 'RU', 'RUSSIAN FEDERATION': 'RU',
              'VIETNAM': 'VN', 'VIET NAM': 'VN', 'TAIWAN': 'TW', 'HONG KONG': 'HK', 'CHINA': 'CN',
              'INDIA': 'IN', 'JAPAN': 'JP', 'AUSTRALIA': 'AU', 'BRAZIL': 'BR', 'CANADA': 'CA',
              'CHILE': 'CL', 'ISRAEL': 'IL', 'MALAYSIA': 'MY', 'MEXICO': 'MX', 'SOUTH AFRICA': 'ZA'}

STATUS_WORDS = r'granted|pending|filed|not in force|withdrawn|lapsed|expired|abandoned|rejected|to be confirmed|term extended'
NO_NATIONAL = re.compile(r'not (yet )?(entered )?(in|into) national phase|national phase not entered|not yet entered in national phase|unknown - national phase', re.I)
FILLER = re.compile(r'\b(the|only|and|patent application is|patent|application|is|in|member states?)\b', re.I)

# --- fetching ----------------------------------------------------------------------------------

def fetch(url, path, refresh):
    if os.path.exists(path) and os.path.getsize(path) and not refresh:
        return open(path, 'rb').read()
    os.makedirs(os.path.dirname(path), exist_ok=True)
    req = urllib.request.Request(url, headers={'User-Agent': 'MPP-LAPaL-prototype-build'})
    for attempt in range(3):
        try:
            data = urllib.request.urlopen(req, timeout=60).read()
            open(path, 'wb').write(data)
            return data
        except Exception as e:  # network hiccup: retry, then give up loudly
            err = e; time.sleep(2)
    raise RuntimeError('could not fetch %s: %s' % (url, err))

def fetch_json(url, path, refresh):
    try:
        return json.loads(fetch(url, path, refresh))
    except Exception as e:
        return {'_error': str(e)}

# --- countries ---------------------------------------------------------------------------------

def load_countries(refresh):
    raw = fetch_json('https://api.worldbank.org/v2/country?format=json&per_page=400',
                     os.path.join(CACHE, 'reference', 'worldbank_countries.json'), refresh)
    countries = {}
    for c in raw[1]:
        if c['region']['value'] == 'Aggregates':
            continue
        inc = c['incomeLevel']['value'].replace(' income', '').replace('Upper middle', 'Upper-middle').replace('Lower middle', 'Lower-middle')
        countries[c['iso2Code']] = {'iso2': c['iso2Code'], 'iso3': c['id'], 'name': c['name'], 'income': inc,
                                    'lmic': inc in ('Low', 'Lower-middle', 'Upper-middle')}
    # Not in the World Bank list but appear in patent data.
    countries.setdefault('TW', {'iso2': 'TW', 'iso3': 'TWN', 'name': 'Taiwan', 'income': 'High', 'lmic': False,
                                'note': 'not in the World Bank list; treated as high income'})
    for iso2, iso3, name in (('AI', 'AIA', 'Anguilla'), ('MS', 'MSR', 'Montserrat')):
        countries.setdefault(iso2, {'iso2': iso2, 'iso3': iso3, 'name': name, 'income': 'Not classified', 'lmic': False,
                                    'note': 'UK overseas territory, not in the World Bank list'})
    return countries

def load_licences(countries, report):
    """Territory of each voluntary licence: LMIC iso2 -> coverage basis ('Direct', 'Indirect', ...)."""
    if not os.path.exists(LICENCE_XLSX):
        report['errors'].append('Licence workbook not found: %s' % LICENCE_XLSX); return []
    wb = openpyxl.load_workbook(LICENCE_XLSX, read_only=True, data_only=True)
    by3 = {c['iso3']: k for k, c in countries.items()}
    out = []
    for lic in LICENCES:
        terr = {}
        for r in wb[lic['sheet']].iter_rows(min_row=5, values_only=True):
            iso2 = by3.get(r[lic['iso3_col']])
            if iso2 and countries[iso2]['lmic'] and r[lic['status_col']] == 'Yes':
                terr[iso2] = (r[lic['basis_col']] if lic['basis_col'] is not None else None) or 'In territory'
        out.append({k: v for k, v in lic.items() if not k.endswith('_col') and k != 'sheet'} | {'territory': terr})
    return out

def name_index(countries):
    idx = {c['name'].upper(): k for k, c in countries.items()}
    idx.update(NAME_ALIAS)
    return idx

# --- inputs ------------------------------------------------------------------------------------

def read_confirmed():
    ws = openpyxl.load_workbook(REVIEW, read_only=True)['Confirm URLs']
    keep, excluded = [], []
    for r in ws.iter_rows(min_row=2, values_only=True):
        row, name, ta, kind, uuid, url = r[0], r[1], r[2], r[3], r[4], r[5]
        confirm = (str(r[11]).strip().upper() if r[11] else '')
        corrected = (str(r[12]).strip() if r[12] else '')
        if not name:
            excluded.append((row, 'row cleared in the review sheet')); continue
        if confirm == 'N':
            if corrected.lower() == 'exclude' or not corrected:
                excluded.append((row, 'marked N' + (' / exclude' if corrected else ' with no corrected URL'))); continue
            url = corrected
        keep.append({'row': row, 'name': str(name).strip(), 'url': url, 'uuid': (uuid or '').strip(),
                     'confirm': confirm or 'blank (accepted)'})
    return keep, excluded

def read_master():
    ws = openpyxl.load_workbook(MASTER, read_only=True)['LAPaL Follow-up']
    return {i: {'category': r[1], 'ta': r[5], 'company': r[4]}
            for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2) if any(r)}

# --- LAPaL page --------------------------------------------------------------------------------

def page_info(raw):
    x = raw.decode('utf-8', 'replace')
    t = re.search(r'<title>(.*?)</title>', x, re.S)
    title = html.unescape(t.group(1)).split('|')[0].strip() if t else ''
    m = re.search(r':medspal="([^"]*)"', x)
    uuid = None
    if m:
        try:
            uuid = json.loads(html.unescape(m.group(1))).get('uuid')
        except Exception:
            pass
    pane = x[x.find('id="patent"'):] if 'id="patent"' in x else ''
    cards = []
    for body in re.findall(r'<div class="card-body">(.*?)</div>\s*</div>\s*</div>', pane, re.S):
        pairs = re.findall(r'<div class="label">(.*?)</div>\s*<p class="card-text">(.*?)</p>', body, re.S)
        d = {k.strip(): html.unescape(re.sub(r'<[^>]+>', ' ', v)).strip() for k, v in pairs}
        if 'Representative patent' in d:
            cards.append(d)
    return title, uuid, cards

# --- card status parsing -----------------------------------------------------------------------

def resolve(token, names, countries):
    """One country-ish token -> (list of iso2, via region or None, how) or None if not understood."""
    t = token.strip().strip('.;:').strip()
    t = re.sub(r'^LMICs:\s*', '', t, flags=re.I)
    if not t:
        return []
    # Bare codes first: IN (India) and IS (Iceland) would otherwise be removed as filler words.
    if re.fullmatch(r'[A-Z]{2}', t) and t in countries and t not in REGION_ALIAS:
        return [(t, None, 'code')]
    if t.upper() in REGION_ALIAS or t.upper() in NAME_ALIAS:  # before filler words are stripped
        up = t.upper()
        if up in NAME_ALIAS:
            return [(NAME_ALIAS[up], None, 'country name')]
        reg = REGION_ALIAS[up]
        return [(c, reg, 'all members of %s (office named alone)' % reg) for c in REGIONS[reg]['members']]
    up = FILLER.sub(' ', t).upper()
    up = re.sub(r'\s+', ' ', up).strip()
    if not up:
        return []
    if up in REGION_ALIAS:
        reg = REGION_ALIAS[up]
        return [(c, reg, 'all members of %s (office named alone)' % reg) for c in REGIONS[reg]['members']]
    if up in CODE_ALIAS:
        return [(CODE_ALIAS[up], None, 'code')]
    if re.fullmatch(r'[A-Z]{2}', up) and up in countries:
        return [(up, None, 'code')]
    if up in names:
        return [(names[up], None, 'country name')]
    parts = up.split()  # "RU US" with a missing comma
    if len(parts) > 1 and all(p in countries or p in CODE_ALIAS for p in parts):
        return [(CODE_ALIAS.get(p, p), None, 'code') for p in parts]
    return None

def split_tokens(seg):
    """Split a country list on commas and 'and', keeping 'EP (AT, BE, ...)' groups together.
    Text right after a closing bracket with no comma ("EP (..., TR,) SG") starts a new token."""
    out, depth, cur, i = [], 0, '', 0
    while i < len(seg):
        ch = seg[i]
        if ch == '(':
            depth += 1
        elif ch == ')':
            depth -= 1
            if depth == 0:
                out.append(cur + ch); cur = ''; i += 1; continue
        if depth == 0 and (ch == ',' or seg[i:i + 5].lower() == ' and '):
            out.append(cur); cur = ''
            i += 5 if ch != ',' else 1
            continue
        cur += ch; i += 1
    out.append(cur)
    # A group token "EP (..." is only complete when its bracket follows the office name.
    merged = []
    for tok in out:
        if tok.strip().startswith('(') and merged:
            merged[-1] += ' ' + tok
        else:
            merged.append(tok)
    return [t.strip() for t in merged if t.strip()]

def parse_status(text, family_expiry, names, countries):
    """Read a LAPaL card status line. Returns countries, family-level status, unparsed bits, method."""
    flat = re.sub(r'\s+', ' ', text.replace('\n', ' / ')).strip()
    result = {'countries': [], 'family_status': None, 'unparsed': [], 'method': None, 'notes': []}
    if not flat or flat.lower() == 'not provided':
        result['method'] = 'no status given'; return result
    if NO_NATIONAL.search(flat):
        result['family_status'] = 'PCT stage, no national filings yet'
        result['method'] = 'no countries'; result['notes'].append(flat); return result
    marks = list(re.finditer(r'(?<![A-Za-z])(%s)(?:\s+in\s*(?:the\s+)?)?\s*:?' % STATUS_WORDS, flat, re.I))
    if not marks:
        result['family_status'] = flat
        result['method'] = 'no countries'
        return result
    structured = ':' in flat
    found_any = False
    for i, m in enumerate(marks):
        status = m.group(1).lower()
        seg = flat[m.end(): marks[i + 1].start() if i + 1 < len(marks) else len(flat)]
        seg = seg.strip(' /-')
        if status == 'term extended':
            for num, until in re.findall(r'\b([A-Z]{2})\d+\w*\s*\(until (\d{2}\.\d{2}\.\d{4})\)', seg):
                exp = dt.datetime.strptime(until, '%d.%m.%Y').date().isoformat()
                result['countries'].append({'iso2': num, 'status_raw': 'Term extended', 'via': None,
                                            'how': 'term extension', 'expiry': exp})
                found_any = True
            continue
        bracket_only = re.fullmatch(r'\(([^()]*)\)', seg)
        if bracket_only:  # "Pending (US only)", "Expired (US, GB)"
            seg = bracket_only.group(1)
        if not seg:
            continue
        for tok in split_tokens(seg):
            g = re.fullmatch(r'([A-Za-z ]+?)\s*\((.*)\)\s*', tok)
            if g and g.group(1).strip().upper() in REGION_ALIAS:  # office with its own member list
                reg = REGION_ALIAS[g.group(1).strip().upper()]
                inner = [x for x in split_tokens(g.group(2)) if x]
                for x in inner:
                    r = resolve(x, names, countries)
                    if r is None:
                        result['unparsed'].append(x)
                    for iso, _, how in (r or []):
                        result['countries'].append({'iso2': iso, 'status_raw': status.capitalize(), 'via': reg,
                                                    'how': 'listed under %s' % reg, 'expiry': family_expiry})
                        found_any = True
                continue
            tok = re.sub(r'\(.*?\)', '', tok).strip()  # e.g. "(US10584098B2)" after a country
            r = resolve(tok, names, countries)
            if r is None:
                result['unparsed'].append(tok); continue
            for iso, via, how in r:
                result['countries'].append({'iso2': iso, 'status_raw': status.capitalize(), 'via': via,
                                            'how': how, 'expiry': family_expiry})
                found_any = True
    if not found_any:  # e.g. "Expired - Fee Expired": a status for the family, not a country list
        result['family_status'] = flat
        result['method'] = 'no countries'
        result['unparsed'] = []
    else:
        result['method'] = 'structured list' if structured else 'free text'
    return result

# --- MedsPaL -----------------------------------------------------------------------------------

def patent_json(puuid, refresh):
    return fetch_json('https://assets.medspal.org/patent/%s.json' % puuid,
                      os.path.join(CACHE, 'medspal', 'patent', puuid + '.json'), refresh)

def medspal_families(uuid, refresh, countries, report):
    """Families and per-country records for one MedsPaL product, built to agree with MedsPaL's
    own report map (medspal.org/reports).

    MedsPaL stores one record per document per country. A record's JSON gives its own status at
    the top level and lists the *other* records of its family under 'related_patents', so the
    representative WO alone never shows itself, nor records only reachable from elsewhere (e.g. some
    EP validations, GCC patents). We therefore merge, by record uuid:
    - the representative patent of each family (product JSON -> patent JSON), and
    - every in-force record the report JSON names (reports/patent/any/<uuid>.json, 'documents':
      country -> record uuid), each with its own related records.
    The report's 'granted'/'filed' country sets are kept to check the result against medspal.org.
    The product 'areas' family summary (what the LAPaL patent map draws) fills in any country where
    MedsPaL has no national record for that family, labelled as family-level (Lobna, 27 Oct). Records
    always win where they exist. Mostly this adds high-income countries; for the LMICs it adds, the map
    can differ from medspal.org's report map, which uses records only.
    """
    base = os.path.join(CACHE, 'medspal')
    prod = fetch_json('https://assets.medspal.org/product/%s.json' % uuid, os.path.join(base, 'product', uuid + '.json'), refresh)
    if '_error' in prod:
        report['errors'].append('MedsPaL product %s: %s' % (uuid, prod['_error'])); return None, [], {}
    rep = {}
    for kind in ('any', 'compound'):
        r = fetch_json('https://assets.medspal.org/reports/patent/%s/%s.json?cache=26' % (kind, uuid),
                       os.path.join(base, 'report_' + kind, uuid + '.json'), refresh)
        rep[kind] = r if '_error' not in r else {'filters': [], 'documents': []}

    records = {}  # record uuid -> record

    def ingest(js):
        if '_error' in js:
            return
        top = {'uuid': js.get('uuid'), 'docdb_number': js.get('number'), 'country_code': js.get('country_code'),
               'status': js.get('status'), 'expiration_date': js.get('expiration_date'),
               'simple_family_uuid': js.get('simple_family_uuid'), 'date': js.get('date')}
        for r in [top] + (js.get('related_patents') or []):
            code = MEDSPAL_CODE_ALIAS.get(r.get('country_code'), r.get('country_code'))
            if not r.get('uuid'):
                continue
            if code in countries:
                records.setdefault(r['uuid'], dict(r, country_code=code))
            elif code in HISTORICAL_CODES:
                report['historical_codes'][code] += 1
            elif code != 'WO':  # the PCT application itself is not a country
                report['unknown_codes'][code] += 1

    fams, by_sf = [], {}
    for f in prod.get('simple_families', []):
        p = f['patent']
        puuid = p['medspal_url'].rstrip('/').split('/')[-1].replace('.json', '')
        pat = patent_json(puuid, refresh)
        if '_error' in pat:
            report['errors'].append('MedsPaL patent %s (%s): %s' % (puuid, p.get('publication_number'), pat['_error']))
        ingest(pat)
        fam_exp = p.get('expiry_date') or None
        filing = pat.get('date') or None
        fam = {'id': f['uuid'], 'description': (f.get('description') or '').strip(),
               'brief_description': (f.get('abstract') or '').strip(),
               'representative_patent': p.get('publication_number'), 'espacenet_url': p.get('espacenet_url'),
               'categories': f.get('categories') or [], 'holders': [re.sub(r'\s+', ' ', a).strip() for a in p.get('applicants') or []],
               'source': f.get('source'), 'expiry': fam_exp,
               'filing_date': filing, 'filing_estimated': False, 'family_status': None, 'countries': []}
        if not filing and fam_exp:
            fam['filing_date'] = shift_years(fam_exp, -20); fam['filing_estimated'] = True
        fams.append(fam); by_sf[f['uuid']] = fam

    for d in rep['any'].get('documents', []):
        for c in d['content']:
            if c['uuid'] not in records or not records[c['uuid']].get('_full'):
                js = patent_json(c['uuid'], refresh)
                if '_error' in js:
                    report['errors'].append('MedsPaL record %s %s (%s): %s' % (c['name'], d['country'], prod.get('name', '').strip(), js['_error']))
                    continue
                ingest(js)
                if c['uuid'] in records:
                    records[c['uuid']]['_full'] = True

    # The report's document -> countries list is the set of records MedsPaL curators linked to this
    # product (expired ones included). Family records not linked to the product are left out, as on
    # medspal.org; they are counted in the build report.
    doc_filters = [f for f in rep['any'].get('filters', []) if f['status'] not in ('granted', 'filed')]
    linked = {(f['status'], MEDSPAL_CODE_ALIAS.get(c, c)) for f in doc_filters for c in f['countries']}
    multi = {f['status'] for f in doc_filters if len(f['countries']) > 1}
    outside = Counter(); unlinked = 0
    for r in records.values():
        if (r.get('docdb_number'), r['country_code']) not in linked:
            unlinked += 1; continue
        fam = by_sf.get(r.get('simple_family_uuid'))
        if fam is None:
            outside[r.get('docdb_number')] += 1; continue
        raw = r.get('status') or ''
        cls = STATUS_CLASS.get(raw.lower())
        if cls is None:
            report['unmapped_status'][raw] += 1
        exp = r.get('expiration_date') or None
        if cls in ('granted', 'filed') and not exp:
            exp = fam['expiry']
        number = r.get('docdb_number')
        fam['countries'].append({'iso2': r['country_code'], 'status_raw': raw, 'status': cls,
                                 'expiry': exp if cls != 'not_in_force' else None, 'number': number,
                                 'via': number[:2] if number in multi else None})
    report['unlinked'].append((prod.get('name', '').strip(), unlinked, len(linked)))
    # MedsPaL's "compound patent" toggle is driven by its own compound report, which does not
    # always follow the family category: a family counts as compound if any of its records is there.
    compound_docs = {f['status'] for f in rep['compound'].get('filters', []) if f['status'] not in ('granted', 'filed')}
    for fam in fams:
        fam['compound_patent_medspal'] = any(c['number'] in compound_docs for c in fam['countries'])
    # Family-level status from the product summary wherever no national record exists (HICs and LMICs).
    code = {'g': 'Granted', 'f': 'Filed', 'n': 'Not in force'}
    for pf in prod.get('simple_families', []):
        fam = by_sf.get(pf['uuid'])
        if not fam:
            continue
        have = {c['iso2'] for c in fam['countries']}
        for a in pf.get('areas', []):
            iso = MEDSPAL_CODE_ALIAS.get(a['iso2'], a['iso2'])
            if iso not in countries or iso in have or not a.get('legal_status'):
                continue
            raw = code.get(a['legal_status'][0])
            if not raw:
                continue
            cls = STATUS_CLASS[raw.lower()]
            if countries[iso]['lmic']:
                report['lmic_summary_only'] += 1
            fam['countries'].append({'iso2': iso, 'status_raw': raw, 'status': cls,
                                     'expiry': fam['expiry'] if cls != 'not_in_force' else None, 'number': None,
                                     'via': None, 'source': 'family summary'})
            report['hic_summary' if not countries[iso]['lmic'] else 'lmic_summary'][cls] += 1
    for n, k in outside.items():
        report['doc_outside_families'].append('%s (%d record(s)) in %s' % (n, k, prod.get('name', '').strip()))
    for fam in fams:
        if not fam['countries']:
            report['no_countries'].append('%s (%s)' % (fam['representative_patent'], prod.get('name', '').strip()))
    medspal_sets = {kind: {f['status']: {MEDSPAL_CODE_ALIAS.get(c, c) for c in f['countries']}
                           for f in rep[kind].get('filters', []) if f['status'] in ('granted', 'filed')}
                    for kind in rep}
    return prod, fams, medspal_sets

# --- filing dates for LAPaL card families (Google Patents) ---------------------------------------

GP_CACHE = os.path.join(CACHE, 'google_patents_filing.json')

def clean_number(raw):
    """Tidy a card's representative patent number for lookup: first number only, letter O typed
    for zeros in 'USOO…', '!' typed for '1', kind code in brackets."""
    first = re.split(r'[,;/]', raw or '')[0].strip()
    first = re.sub(r'\s*\(([A-Z]\d?)\)$', r'\1', first).replace('!', '1').replace(' ', '')
    first = re.sub(r'^USOO', 'US', first)
    return first if re.fullmatch(r'[A-Z]{2}\d+[A-Z]?\d?', first) else None

def google_filing(raw, refresh, cache, report):
    """Filing date of the representative patent from its public Google Patents page
    (meta DC.date scheme=dateSubmitted: the international filing date for a WO publication, as
    MedsPaL uses; the application filing date otherwise). Only the dates are cached."""
    num = clean_number(raw)
    if not num:
        report['filing_lookup'].append('%s: number not recognised' % raw); return None
    if num in cache and not refresh:
        return cache[num]
    hit = None
    for candidate in [num] + ([num + k for k in ('A1', 'B2', 'B1', 'A')] if not re.search(r'[A-Z]\d?$', num[2:]) else []):
        try:
            req = urllib.request.Request('https://patents.google.com/patent/%s/en' % candidate,
                                         headers={'User-Agent': 'Mozilla/5.0 (MPP LAPaL prototype build)'})
            page = urllib.request.urlopen(req, timeout=60).read().decode('utf-8', 'replace')
        except Exception:
            time.sleep(1); continue
        filed = re.search(r'<meta name="DC.date" content="(\d{4}-\d{2}-\d{2})" scheme="dateSubmitted">', page)
        prio = re.search(r'itemprop="priorityDate" datetime="(\d{4}-\d{2}-\d{2})"', page)
        if filed:
            hit = {'filing': filed.group(1), 'priority': prio.group(1) if prio else None,
                   'url': 'https://patents.google.com/patent/%s/en' % candidate}
            break
        time.sleep(1)
    time.sleep(1)  # be polite to Google Patents
    cache[num] = hit
    json.dump(cache, open(GP_CACHE, 'w'), indent=1, sort_keys=True)
    if not hit:
        report['filing_lookup'].append('%s: no filing date found on Google Patents' % raw)
    return hit

def shift_years(iso, years):
    d = dt.date.fromisoformat(iso)
    try:
        return d.replace(year=d.year + years).isoformat()
    except ValueError:  # 29 Feb
        return d.replace(year=d.year + years, day=28).isoformat()

# --- main --------------------------------------------------------------------------------------

def collapse(countries_list):
    """Several records for one country in one family (divisionals): keep them all, add the
    strongest as the country's status for the map."""
    by = defaultdict(list)
    for c in countries_list:
        by[c['iso2']].append(c)
    out = []
    for iso, recs in sorted(by.items()):
        mapped = [r for r in recs if r.get('status')]
        best = max(mapped, key=lambda r: RANK[r['status']]) if mapped else recs[0]
        exps = [r['expiry'] for r in recs if r.get('expiry') and r.get('status') in ('granted', 'filed')]
        out.append({'iso2': iso, 'status': best.get('status'), 'status_raw': best['status_raw'],
                    'expiry': max(exps) if exps else best.get('expiry'),
                    'records': [{k: v for k, v in r.items() if k != 'iso2'} for r in recs]})
    return out

def main():
    ap = argparse.ArgumentParser(); ap.add_argument('--refresh', action='store_true')
    args = ap.parse_args()
    os.makedirs(OUT, exist_ok=True)
    countries = load_countries(args.refresh); names = name_index(countries)
    rows, excluded = read_confirmed(); master = read_master()
    report = {'errors': [], 'unmapped_status': Counter(), 'doc_outside_families': [], 'no_countries': [],
              'uuid_flags': [], 'unknown_codes': Counter(), 'medspal_check': [], 'historical_codes': Counter(), 'unlinked': [],
              'filing_lookup': [], 'hic_summary': Counter(), 'lmic_summary': Counter(), 'lmic_summary_only': 0, 'hic_dropped': 0}
    gp_cache = json.load(open(GP_CACHE)) if os.path.exists(GP_CACHE) else {}
    card_review = []
    licences = load_licences(countries, report)
    report['licences'] = licences

    by_url = defaultdict(list)
    for r in rows:
        by_url[r['url']].append(r)

    entries = []
    for url, rs in sorted(by_url.items()):
        slug = url.rstrip('/').split('/')[-1]; section = url.rstrip('/').split('/')[-2]
        raw = fetch(url, os.path.join(CACHE, 'lapal', section + '__' + slug + '.html'), args.refresh)
        title, page_uuid, cards = page_info(raw)
        # Split therapeutic areas on commas/semicolons, but not inside brackets ("Pain (incl. joints, knee)").
        tas = sorted({t.strip() for r in rs for t in re.split(r'[,;](?![^()]*\))', str(master.get(r['row'], {}).get('ta') or '')) if t.strip()})
        row_uuids = sorted({r['uuid'] for r in rs if r['uuid']})
        entry = {'id': section + '/' + slug, 'url': url, 'name': title or rs[0]['name'],
                 'type': {'compounds': 'Compound', 'formulations': 'Formulation & regimen', 'technology': 'Technology'}.get(section, section),
                 'therapeutic_areas': tas, 'hiv': any('hiv' in t.lower() for t in tas),
                 'workbook_rows': [r['row'] for r in rs], 'workbook_names': [r['name'] for r in rs]}
        uuid = page_uuid or (row_uuids[0] if row_uuids else None)
        if page_uuid and row_uuids and page_uuid not in row_uuids:
            report['uuid_flags'].append('%s: page uses MedsPaL %s, workbook row(s) %s give %s; page uuid used'
                                        % (entry['name'], page_uuid, entry['workbook_rows'], ', '.join(row_uuids)))
        if not page_uuid and row_uuids:
            report['uuid_flags'].append('%s: no MedsPaL uuid on the LAPaL page; workbook uuid %s used'
                                        % (entry['name'], row_uuids[0]))
        if uuid:
            prod, fams, medspal_sets = medspal_families(uuid, args.refresh, countries, report)
            entry.update({'source': 'medspal', 'preliminary': False, 'medspal_uuid': uuid,
                          'medspal_product_name': (prod or {}).get('name', '').strip() if prod else None})
        else:
            fams = []
            for i, c in enumerate(cards):
                exp = None
                try:
                    exp = dt.datetime.strptime(c.get('Expiration date', ''), '%B %d, %Y').date().isoformat()
                except ValueError:
                    pass
                st = parse_status(c.get('Status', ''), exp, names, countries)
                gp = google_filing(c.get('Representative patent', ''), args.refresh, gp_cache, report)
                if gp:
                    filing, filing_est, filing_src = gp['filing'], False, gp['url']
                else:  # lookup failed: fall back to the estimate, flagged
                    filing, filing_est, filing_src = (shift_years(exp, -20) if exp else None), bool(exp), None
                cs = []
                for x in st['countries']:
                    cls = STATUS_CLASS.get(x['status_raw'].lower())
                    if x['iso2'] not in countries:
                        report['unknown_codes'][x['iso2']] += 1
                    cs.append(dict(x, status=cls, expiry=x['expiry'] if cls != 'not_in_force' else None))
                fams.append({'id': '%s#%d' % (entry['id'], i + 1), 'description': c.get('Description', ''),
                             'brief_description': c.get('Brief description', ''),
                             'representative_patent': c.get('Representative patent', ''), 'espacenet_url': None,
                             'categories': [x.strip() for x in re.split(r'[,;]', c.get('Category', '')) if x.strip()],
                             'holders': [c.get('Patent holder', '')] if c.get('Patent holder') else [],
                             'exclusivity': c.get('Exclusivity'), 'source': 'LAPaL card', 'expiry': exp,
                             'expiry_raw': c.get('Expiration date', ''),
                             'filing_date': filing, 'filing_estimated': filing_est, 'filing_source': filing_src,
                             'priority_date': gp['priority'] if gp else None,
                             'family_status': st['family_status'], 'status_text': c.get('Status', ''),
                             'parse_method': st['method'], 'unparsed': st['unparsed'], 'countries': cs})
                card_review.append([entry['name'], url, c.get('Representative patent', ''), c.get('Expiration date', ''),
                                    c.get('Status', '').replace('\n', ' / '), st['method'],
                                    summarise(cs), st['family_status'] or '', ', '.join(st['unparsed'])])
            entry.update({'source': 'lapal_card', 'preliminary': True, 'medspal_uuid': None})
        for f in fams:
            # LMICs only (Lobna, 27 Oct): high-income countries are left off the views.
            report['hic_dropped'] += sum(1 for c in f['countries'] if not countries.get(c['iso2'], {}).get('lmic'))
            f['countries'] = collapse([c for c in f['countries'] if countries.get(c['iso2'], {}).get('lmic')])
            f['compound_patent'] = f.pop('compound_patent_medspal', None)
            if f['compound_patent'] is None:  # LAPaL cards: from the card category
                f['compound_patent'] = any(c.lower() == 'compound' for c in f['categories'])
        if uuid:
            lm = {k: {s: {i for i in v if countries.get(i, {}).get('lmic')} for s, v in d.items()} for k, d in medspal_sets.items()}
            report['medspal_check'].append(check_against_medspal(entry['name'], fams, lm))
        fams.sort(key=lambda f: (f['filing_date'] or '9999', f['representative_patent'] or ''))
        entry['families'] = fams
        entries.append(entry)

    snapshot = {
        'meta': {
            'generated_at': dt.datetime.now().isoformat(timespec='seconds'),
            'title': 'LAPaL patent coverage — prototype data snapshot (internal MPP use)',
            'sources': {'lapal': 'https://lapal.ch entry pages (patent cards, embedded MedsPaL uuid)',
                        'medspal': 'https://assets.medspal.org product and patent JSON',
                        'income': 'World Bank country API, current classification'},
            'status_classes': {'granted': 'Granted', 'filed': 'Filed', 'not_in_force': 'Not in force'},
            'status_mapping': STATUS_CLASS,
            'regions': REGIONS,
            'preliminary_disclaimer': 'Patent data for this entry are preliminary. Full country-level coverage from MedsPaL is being added soon.',
            'notes': ['Filing dates for LAPaL-card families come from Google Patents (filing_source); where the lookup failed they are estimated as expiry minus 20 years (filing_estimated = true).',
                      'For a country with several records in one family (e.g. divisionals), status is the strongest of them; all records are kept.',
                      'A regional office named without its member list is expanded to all its members.',
                      'Low- and middle-income countries only: high-income countries are left out of the views.'],
        },
        'countries': countries,
        'licences': licences,
        'entries': entries,
    }
    json.dump(snapshot, open(os.path.join(OUT, 'patents_data.json'), 'w'), ensure_ascii=False, indent=1)
    write_card_review(card_review)
    write_report(entries, excluded, rows, report, countries)
    print(open(os.path.join(OUT, 'build_report.md')).read().split('\n## ')[0])

def status_on(record, day=None):
    """Map class of one record. Today (day=None) the MedsPaL status is used as recorded, even past
    the expiry date, so the map matches medspal.org (Lobna, 26 Oct). For a future date chosen on
    the slider, a granted/filed record whose expiry is before that date is not in force.
    The prototype applies the same rule."""
    st = record.get('status')
    if day and st in ('granted', 'filed') and record.get('expiry') and record['expiry'] < day:
        return 'not_in_force'
    return st

def product_sets(fams, compound_only=False, day=None):
    best = {}
    for f in fams:
        if compound_only and not f['compound_patent']:
            continue
        for c in f['countries']:
            for r in c['records']:
                if r.get('source') == 'family summary':  # not on medspal.org's report map
                    continue
                st = status_on(r, day)
                if st and RANK[st] > RANK.get(best.get(c['iso2']), 0):
                    best[c['iso2']] = st
    return {s: {k for k, v in best.items() if v == s} for s in ('granted', 'filed')}

def check_against_medspal(name, fams, medspal_sets):
    """Our granted/filed countries per product vs the sets MedsPaL's report map uses."""
    out = {'entry': name, 'diffs': []}
    for kind, compound in (('any', False), ('compound', True)):
        ours = product_sets(fams, compound)
        theirs = medspal_sets.get(kind, {})
        for s in ('granted', 'filed'):
            a, b = ours[s], theirs.get(s, set())
            if a != b:
                out['diffs'].append('%s/%s: only ours %s; only MedsPaL %s' % (kind, s, sorted(a - b) or '-', sorted(b - a) or '-'))
    return out

def summarise(cs):
    groups = defaultdict(list)
    for c in cs:
        groups[c['status_raw']].append(c['iso2'] + ('*' if c['via'] and 'alone' in c.get('how', '') else ''))
    return ' / '.join('%s: %s' % (k, ', '.join(v)) for k, v in groups.items())

def write_card_review(rows):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Card status lines'
    ws.append(['Entry', 'LAPaL URL', 'Representative patent', 'Expiration date (card)', 'Status line on the card',
               'How it was read', 'Countries read (* = office named alone, expanded to all members)',
               'Family-level status (no countries)', 'Bits not understood', 'OK? (Y/N)', 'Your correction'])
    grey = PatternFill('solid', fgColor='D9D9D9'); yellow = PatternFill('solid', fgColor='FFF2CC'); blue = PatternFill('solid', fgColor='DDEBF7')
    for r in rows:
        ws.append(r + ['', ''])
    for c in ws[1]:
        c.font = Font(bold=True); c.fill = grey; c.alignment = Alignment(wrap_text=True, vertical='top')
    for row in ws.iter_rows(min_row=2):
        flag = row[5].value in ('free text', 'no countries') or row[8].value or not row[3].value
        for c in row[:9]:
            c.alignment = Alignment(wrap_text=True, vertical='top')
            if flag: c.fill = yellow
        for c in row[9:]: c.fill = blue
    for col, w in zip('ABCDEFGHIJK', [30, 40, 18, 16, 60, 14, 60, 30, 25, 9, 30]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'B2'; ws.auto_filter.ref = ws.dimensions
    wb.save(os.path.join(OUT, 'card_status_review.xlsx'))

def write_report(entries, excluded, rows, rep, countries):
    med = [e for e in entries if e['source'] == 'medspal']; card = [e for e in entries if e['source'] == 'lapal_card']
    fams = [f for e in entries for f in e['families']]
    cf = [f for e in card for f in e['families']]
    L = ['# Build report — LAPaL patent coverage snapshot', '',
         'Generated %s. Snapshot: `data/patents_data.json`.' % dt.datetime.now().strftime('%d %b %Y %H:%M'), '',
         '- **%d entries** from %d confirmed workbook rows (%d excluded)' % (len(entries), len(rows), len(excluded)),
         '- **%d with MedsPaL data**, %d from LAPaL cards (marked preliminary)' % (len(med), len(card)),
         '- **%d HIV entries**, of which %d with MedsPaL data (the set for the "by country" view)' % (
             sum(e['hiv'] for e in entries), sum(e['hiv'] for e in med)),
         '- %d patent families; %d with a real filing date, %d estimated (expiry − 20 y), %d with no expiry date' % (
             len(fams), sum(1 for f in fams if f['filing_date'] and not f['filing_estimated']),
             sum(1 for f in fams if f['filing_estimated']), sum(1 for f in fams if not f['expiry'])),
         '- Card families: %d read as a structured list, %d from free text, %d with no countries, %d with no status' % (
             sum(f['parse_method'] == 'structured list' for f in cf), sum(f['parse_method'] == 'free text' for f in cf),
             sum(f['parse_method'] == 'no countries' for f in cf), sum(f['parse_method'] == 'no status given' for f in cf)),
         '', '## Excluded rows', '']
    L += ['- Follow-up row %s: %s' % x for x in excluded] or ['None']
    L += ['', '## Rows merged into one entry (same LAPaL page)', '']
    L += ['- %s ← rows %s (%s)' % (e['name'], e['workbook_rows'], '; '.join(n.replace('\n', ' ') for n in e['workbook_names']))
          for e in entries if len(e['workbook_rows']) > 1] or ['None']
    L += ['', '## MedsPaL uuid choices', ''] + (['- ' + x for x in rep['uuid_flags']] or ['None'])
    L += ['', '## Low- and middle-income countries only', '',
          'High-income and unclassified countries are left out (Lobna, 27 Oct): %d family × country records dropped. The check below compares LMICs only.' % rep['hic_dropped']]
    ok = [c for c in rep['medspal_check'] if not c['diffs']]
    L += ['', '## Check against the MedsPaL report map', '',
          '%d of %d MedsPaL entries give exactly the same granted and filed LMICs as medspal.org today, for both "any patent" and "compound patent". Remaining differences: "Rejected - under appeal" is not in force here but filed on medspal.org, and a few regional EA / ARIPO / EP records.' % (len(ok), len(rep['medspal_check'])), '']
    L += ['- **%s**: %s' % (c['entry'], ' | '.join(c['diffs'])) for c in rep['medspal_check'] if c['diffs']]
    L += ['', 'Countries with no MedsPaL record for a family use the family-level status from the summary in the product JSON (as the LAPaL patent map does). These are excluded from the check above, which compares MedsPaL records only.',
          '',
          '- LMICs: %d granted, %d filed, %d not in force (family × country). For these, the map can differ from medspal.org.' % (
              rep['lmic_summary']['granted'], rep['lmic_summary']['filed'], rep['lmic_summary']['not_in_force'])]
    L += ['', 'Family records that MedsPaL has not linked to the product are left out, as on medspal.org: %d records across %d products.' % (
        sum(u[1] for u in rep['unlinked']), sum(1 for u in rep['unlinked'] if u[1]))]
    L += ['', '## MedsPaL families with no countries (timeline only)', ''] + (['- ' + x for x in rep['no_countries']] or ['None'])
    L += ['', '## MedsPaL documents not in any of the product\'s families', ''] + (['- ' + x for x in rep['doc_outside_families']] or ['None'])
    L += ['', '## Statuses not mapped to a colour', '']
    L += ['- %s: %d records' % kv for kv in rep['unmapped_status'].items()] or ['None']
    unk = [f for f in cf if any(c['status'] is None for c in f['countries'])]
    L += ['- card %s (%s): %s' % (f['representative_patent'], f['id'], ', '.join('%s %s' % (c['iso2'], c['status_raw']) for c in f['countries'] if c['status'] is None)) for f in unk]
    L += ['', '## Country codes not in the World Bank list', '']
    L += ['- %s (%d)' % kv for kv in rep['unknown_codes'].items()] or ['None']
    L += ['', 'Former states on old MedsPaL records, left off the map: ' +
          (', '.join('%s (%d)' % kv for kv in rep['historical_codes'].items()) or 'none')]
    L += ['', '## Voluntary licences', '', 'Territories read from `licence-territory-dashboard/data/HIV_PrEP_licence_territories_v2_3.xlsx` (LMICs only).', '']
    ids = {e['id'] for e in entries}
    L += ['- **%s** (%s, %s): %d LMICs in territory; entries: %s%s' % (
        l['label'], l['licensor'], l['type'], len(l['territory']), ', '.join(l['entries']),
        '' if all(i in ids for i in l['entries']) else ' — **missing: %s**' % ', '.join(i for i in l['entries'] if i not in ids))
        for l in rep.get('licences', [])] or ['None']
    L += ['', '## Card filing dates (Google Patents)', '',
          '%d of %d card families have the real filing date of their representative patent from Google Patents; %d still estimated (expiry − 20 y), %d with no date.' % (
              sum(1 for f in cf if f.get('filing_source')), len(cf), sum(1 for f in cf if f['filing_estimated']), sum(1 for f in cf if not f['filing_date'])), '']
    L += ['- ' + x for x in rep['filing_lookup']] or ['All lookups succeeded.']
    # A patent normally runs 20 years from filing. A gap well outside that points either to a
    # representative patent filed later than the family (e.g. a US continuation, which keeps the
    # parent's expiry) or to an expiry date on the card that needs checking.
    odd = [f for f in cf if f.get('filing_source') and f['expiry'] and not (18 <= (dt.date.fromisoformat(f['expiry']) - dt.date.fromisoformat(f['filing_date'])).days / 365.25 <= 22)]
    L += ['', '### Filing date and card expiry not about 20 years apart (%d) — please check the card' % len(odd), '',
          '| Card family | Representative patent | Filed (Google Patents) | Expiry on card | Years |', '|---|---|---|---|---|']
    L += ['| %s | %s | %s | %s | %.1f |' % (f['id'], f['representative_patent'], f['filing_date'], f['expiry'],
          (dt.date.fromisoformat(f['expiry']) - dt.date.fromisoformat(f['filing_date'])).days / 365.25) for f in odd]
    L += ['', '## Card families with no expiry date on the card', '']
    L += ['- %s — %s' % (f['id'], f['representative_patent']) for f in cf if not f['expiry']] or ['None']
    L += ['', '## Card status text not understood', '']
    L += ['- %s — %s: %s' % (f['id'], f['representative_patent'], ', '.join(f['unparsed'])) for f in cf if f['unparsed']] or ['None']
    L += ['', '## Errors', ''] + (['- ' + x for x in rep['errors']] or ['None'])
    L += ['', '## Entries', '', '| Entry | Type | TA | Source | Families | Countries (granted / filed / not in force) |', '|---|---|---|---|---|---|']
    for e in entries:
        best = {}
        for f in e['families']:
            for c in f['countries']:
                if c['status'] and RANK[c['status']] > RANK.get(best.get(c['iso2']), 0):
                    best[c['iso2']] = c['status']
        n = Counter(best.values())
        L.append('| [%s](%s) | %s | %s | %s | %d | %d / %d / %d |' % (
            e['name'], e['url'], e['type'], ', '.join(e['therapeutic_areas']),
            'MedsPaL' if e['source'] == 'medspal' else 'LAPaL card (preliminary)', len(e['families']),
            n['granted'], n['filed'], n['not_in_force']))
    open(os.path.join(OUT, 'build_report.md'), 'w').write('\n'.join(L) + '\n')

if __name__ == '__main__':
    main()
